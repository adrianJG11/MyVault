import json
from datetime import date
from decimal import Decimal, InvalidOperation
from hashlib import sha256
from pathlib import Path
from typing import BinaryIO
from zipfile import BadZipFile

from openpyxl import load_workbook


def _parse_ibercaja_date(value: str) -> date:
    day, month, year = value.split("-")
    return date(int(year), int(month), int(day))


def ibercaja_transaction_fingerprint(
    operation_date: date,
    value_date: date,
    amount: Decimal,
    balance_after: Decimal,
    bank_concept: str,
    description: str,
) -> str:
    canonical_values = json.dumps(
        [
            operation_date.isoformat(),
            value_date.isoformat(),
            format(amount, ".2f"),
            format(balance_after, ".2f"),
            bank_concept.strip(),
            description.strip(),
        ],
        ensure_ascii=False,
        separators=(",", ":"),
    )
    return sha256(
        canonical_values.encode("utf-8"),
        usedforsecurity=False,
    ).hexdigest()


def parse_ibercaja_xlsx(
    source: Path | BinaryIO,
) -> list[dict[str, date | Decimal | str]]:
    try:
        workbook = load_workbook(source, read_only=True, data_only=True)
    except BadZipFile as error:
        raise ValueError("Invalid XLSX file") from error
    try:
        sheet = workbook.active

        expected_headers = (
            "Nº Orden",
            "Fecha Oper",
            "Fecha Valor",
            "Concepto",
            "Descripción",
            "Referencia",
            "Importe",
            "Saldo",
        )

        header_row = None

        for row_number, row in enumerate(
            sheet.iter_rows(values_only=True),
            start=1,
        ):
            if row[: len(expected_headers)] == expected_headers:
                header_row = row_number
                break
        if header_row is None:
            raise ValueError("Ibercaja transaction headers were not found")

        transactions: list[dict[str, date | Decimal | str]] = []

        for row_number, row in enumerate(
            sheet.iter_rows(
                min_row=header_row + 1,
                values_only=True,
            ),
            start=header_row + 1,
        ):
            if not isinstance(row[0], int):
                continue

            (
                _order_number,
                operation_date_text,
                value_date_text,
                bank_concept,
                description,
                _reference,
                amount,
                balance_after,
            ) = row[:8]

            if not isinstance(bank_concept, str) or not isinstance(description, str):
                raise ValueError(  # noqa: TRY004
                    f"Invalid Ibercaja transaction row {row_number}"
                )

            try:
                operation_date = _parse_ibercaja_date(operation_date_text)
                value_date = _parse_ibercaja_date(value_date_text)
                decimal_amount = Decimal(str(amount))
                decimal_balance = Decimal(str(balance_after))
            except (AttributeError, InvalidOperation, TypeError, ValueError) as error:
                raise ValueError(
                    f"Invalid Ibercaja transaction row {row_number}"
                ) from error

            transactions.append(
                {
                    "import_fingerprint": ibercaja_transaction_fingerprint(
                        operation_date=operation_date,
                        value_date=value_date,
                        amount=decimal_amount,
                        balance_after=decimal_balance,
                        bank_concept=bank_concept,
                        description=description,
                    ),
                    "operation_date": operation_date,
                    "value_date": value_date,
                    "bank_concept": bank_concept.strip(),
                    "description": description.strip(),
                    "amount": decimal_amount,
                    "balance_after": decimal_balance,
                }
            )

        return transactions
    finally:
        workbook.close()
