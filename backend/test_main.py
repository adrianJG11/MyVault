from collections.abc import Iterator
from datetime import date
from decimal import Decimal
from io import BytesIO

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook
from sqlalchemy import delete

from database import SessionFactory, engine
from investments import router as investments_router
from investments.market_prices import MarketPriceResult
from main import app
from models import Account, InvestmentActivity, InvestmentPrice, Transaction

client = TestClient(app)


def _create_fake_ibercaja_xlsx() -> bytes:
    workbook = Workbook()
    sheet = workbook.active

    headers = [
        "Nº Orden",
        "Fecha Oper",
        "Fecha Valor",
        "Concepto",
        "Descripción",
        "Referencia",
        "Importe",
        "Saldo",
    ]
    for column, header in enumerate(headers, start=1):
        sheet.cell(row=6, column=column, value=header)

    fake_row = [
        1,
        "17-08-2026",
        "17-08-2026",
        "CARD",
        "MERCADONA CÑ DIEGO",
        "FAKE123",
        -42.64,
        14897.84,
    ]
    for column, value in enumerate(fake_row, start=1):
        sheet.cell(row=7, column=column, value=value)

    with BytesIO() as workbook_file:
        workbook.save(workbook_file)
        workbook.close()
        return workbook_file.getvalue()


def _create_fake_revolut_investment_csv() -> bytes:
    return (
        b"Date,Ticker,Type,Quantity,Price per share,Total Amount,Currency,FX Rate\n"
        b"2026-01-10T12:30:00Z,FAKE,BUY - MARKET,0.5,USD 100,USD 50,USD,1.2\n"
        b"2026-02-01T09:00:00Z,FAKE,DIVIDEND,,,USD 0.25,USD,1.1\n"
    )


@pytest.fixture
def clean_database() -> Iterator[None]:
    assert engine.url.database == "finanzas_test"

    with SessionFactory.begin() as session:
        session.execute(delete(InvestmentPrice))
        session.execute(delete(InvestmentActivity))
        session.execute(delete(Transaction))
        session.execute(delete(Account))

    yield

    with SessionFactory.begin() as session:
        session.execute(delete(InvestmentPrice))
        session.execute(delete(InvestmentActivity))
        session.execute(delete(Transaction))
        session.execute(delete(Account))


def test_health() -> None:
    response = client.get("/health")

    assert response.status_code == 200
    assert response.json() == {"status": "ok"}


def test_create_account_returns_account(clean_database: None) -> None:
    account_data = {
        "name": "Main account",
        "bank_name": "IGN",
        "currency": "EUR",
    }
    response = client.post("/accounts", json=account_data)
    response_data = response.json()

    assert response.status_code == 201
    assert isinstance(response_data["id"], int)
    assert response_data["id"] > 0
    assert response_data == {
        "id": response_data["id"],
        **account_data,
    }


def test_create_account_rejects_missing_currency() -> None:
    incomplete_account = {
        "name": "Main account",
        "bank_name": "ING",
    }

    response = client.post("/accounts", json=incomplete_account)

    assert response.status_code == 422


def test_list_accounts_returns_created_account(clean_database: None) -> None:
    account_data = {
        "name": "Main account",
        "bank_name": "IGN",
        "currency": "EUR",
    }

    create_response = client.post("/accounts", json=account_data)

    assert create_response.status_code == 201

    created_account = create_response.json()
    list_response = client.get("/accounts")

    assert list_response.status_code == 200

    accounts = list_response.json()
    assert accounts == [created_account]


def test_list_empty_database(clean_database: None) -> None:
    list_response = client.get("/accounts")
    assert list_response.status_code == 200
    accounts = list_response.json()
    assert accounts == []


def test_list_transactions_returns_empty_list(clean_database: None) -> None:
    list_response = client.get("/transactions")
    assert list_response.status_code == 200
    transactions = list_response.json()
    assert transactions == []


def test_list_transactions_returns_stored_transaction(clean_database: None) -> None:
    account_data = {
        "name": "Main account",
        "bank_name": "ING",
        "currency": "EUR",
    }

    create_response = client.post("/accounts", json=account_data)

    assert create_response.status_code == 201

    account_id = create_response.json()["id"]

    with SessionFactory.begin() as session:
        transaction = Transaction(
            account_id=account_id,
            import_fingerprint="a" * 64,
            operation_date=date(2026, 8, 17),
            value_date=date(2026, 8, 17),
            amount=Decimal("-42.64"),
            balance_after=Decimal("14897.84"),
            bank_concept="CARD",
            description="FAKE SUPERMARKET",
        )
        session.add(transaction)

    list_response = client.get("/transactions")

    assert list_response.status_code == 200

    transactions = list_response.json()
    assert len(transactions) == 1

    transaction_data = transactions[0]
    assert isinstance(transaction_data["id"], int)
    assert transaction_data == {
        "id": transaction_data["id"],
        "account_id": account_id,
        "operation_date": "2026-08-17",
        "value_date": "2026-08-17",
        "amount": "-42.64",
        "balance_after": "14897.84",
        "bank_concept": "CARD",
        "description": "FAKE SUPERMARKET",
        "category": None,
    }


def test_list_transactions_filters_by_account_id(clean_database: None) -> None:
    first_account_response = client.post(
        "/accounts",
        json={
            "name": "First account",
            "bank_name": "Ibercaja",
            "currency": "EUR",
        },
    )
    second_account_response = client.post(
        "/accounts",
        json={
            "name": "Second account",
            "bank_name": "Other bank",
            "currency": "EUR",
        },
    )

    assert first_account_response.status_code == 201
    assert second_account_response.status_code == 201

    first_account_id = first_account_response.json()["id"]
    second_account_id = second_account_response.json()["id"]

    with SessionFactory.begin() as session:
        session.add_all(
            [
                Transaction(
                    account_id=first_account_id,
                    import_fingerprint="b" * 64,
                    operation_date=date(2026, 8, 17),
                    value_date=date(2026, 8, 17),
                    amount=Decimal("-42.64"),
                    balance_after=Decimal("14897.84"),
                    bank_concept="CARD",
                    description="FIRST ACCOUNT PURCHASE",
                ),
                Transaction(
                    account_id=second_account_id,
                    import_fingerprint="c" * 64,
                    operation_date=date(2026, 8, 17),
                    value_date=date(2026, 8, 17),
                    amount=Decimal("-18.57"),
                    balance_after=Decimal("1000.00"),
                    bank_concept="CARD",
                    description="SECOND ACCOUNT PURCHASE",
                ),
            ]
        )

    response = client.get(
        "/transactions",
        params={"account_id": first_account_id},
    )

    assert response.status_code == 200

    transactions = response.json()
    assert len(transactions) == 1
    assert transactions[0]["account_id"] == first_account_id
    assert transactions[0]["description"] == "FIRST ACCOUNT PURCHASE"


def test_list_transactions_filters_by_operation_date_range(
    clean_database: None,
) -> None:
    create_response = client.post(
        "/accounts",
        json={
            "name": "Main account",
            "bank_name": "Ibercaja",
            "currency": "EUR",
        },
    )

    assert create_response.status_code == 201

    account_id = create_response.json()["id"]
    dated_transactions = [
        (date(2026, 8, 9), "BEFORE RANGE"),
        (date(2026, 8, 10), "START BOUNDARY"),
        (date(2026, 8, 20), "END BOUNDARY"),
        (date(2026, 8, 21), "AFTER RANGE"),
    ]

    with SessionFactory.begin() as session:
        for index, (operation_date, description) in enumerate(dated_transactions):
            session.add(
                Transaction(
                    account_id=account_id,
                    import_fingerprint=f"{index:064x}",
                    operation_date=operation_date,
                    value_date=operation_date,
                    amount=Decimal("-1.00"),
                    balance_after=Decimal("1000.00") - index,
                    bank_concept="TEST",
                    description=description,
                )
            )

    response = client.get(
        "/transactions",
        params={
            "date_from": "2026-08-10",
            "date_to": "2026-08-20",
        },
    )

    assert response.status_code == 200

    transactions = response.json()
    assert len(transactions) == 2
    assert {transaction["description"] for transaction in transactions} == {
        "START BOUNDARY",
        "END BOUNDARY",
    }


def test_list_transactions_rejects_reversed_date_range(
    clean_database: None,
) -> None:
    response = client.get(
        "/transactions",
        params={
            "date_from": "2026-08-20",
            "date_to": "2026-08-10",
        },
    )

    assert response.status_code == 422
    assert response.json() == {"detail": "date_from must be before or equal to date_to"}


def test_list_transactions_orders_newest_first(clean_database: None) -> None:
    create_response = client.post(
        "/accounts",
        json={
            "name": "Main account",
            "bank_name": "Ibercaja",
            "currency": "EUR",
        },
    )

    assert create_response.status_code == 201

    account_id = create_response.json()["id"]
    dated_transactions = [
        (date(2026, 8, 10), "OLDEST"),
        (date(2026, 8, 20), "NEWEST FIRST"),
        (date(2026, 8, 20), "NEWEST SECOND"),
    ]

    with SessionFactory.begin() as session:
        for index, (operation_date, description) in enumerate(dated_transactions):
            session.add(
                Transaction(
                    account_id=account_id,
                    import_fingerprint=f"{index + 20:064x}",
                    operation_date=operation_date,
                    value_date=operation_date,
                    amount=Decimal("-1.00"),
                    balance_after=Decimal("1000.00") - index,
                    bank_concept="TEST",
                    description=description,
                )
            )

    response = client.get("/transactions")

    assert response.status_code == 200
    assert [transaction["description"] for transaction in response.json()] == [
        "NEWEST SECOND",
        "NEWEST FIRST",
        "OLDEST",
    ]


def test_list_transactions_filters_by_description(clean_database: None) -> None:
    create_response = client.post(
        "/accounts",
        json={
            "name": "Main account",
            "bank_name": "Ibercaja",
            "currency": "EUR",
        },
    )

    assert create_response.status_code == 201

    account_id = create_response.json()["id"]

    with SessionFactory.begin() as session:
        session.add_all(
            [
                Transaction(
                    account_id=account_id,
                    import_fingerprint=f"{30:064x}",
                    operation_date=date(2026, 8, 17),
                    value_date=date(2026, 8, 17),
                    amount=Decimal("-42.64"),
                    balance_after=Decimal("14897.84"),
                    bank_concept="CARD",
                    description="MERCADONA CÑ DIEGO",
                ),
                Transaction(
                    account_id=account_id,
                    import_fingerprint=f"{31:064x}",
                    operation_date=date(2026, 8, 16),
                    value_date=date(2026, 8, 16),
                    amount=Decimal("2000.00"),
                    balance_after=Decimal("14940.48"),
                    bank_concept="TRANSFER",
                    description="SALARY",
                ),
            ]
        )

    response = client.get(
        "/transactions",
        params={"description": "mercadona"},
    )

    assert response.status_code == 200

    transactions = response.json()
    assert len(transactions) == 1
    assert transactions[0]["description"] == "MERCADONA CÑ DIEGO"


def test_list_transactions_filters_by_amount_range(clean_database: None) -> None:
    create_response = client.post(
        "/accounts",
        json={
            "name": "Main account",
            "bank_name": "Ibercaja",
            "currency": "EUR",
        },
    )

    assert create_response.status_code == 201

    account_id = create_response.json()["id"]
    transactions_by_amount = [
        (Decimal("-100.00"), "BELOW RANGE"),
        (Decimal("-50.00"), "MINIMUM BOUNDARY"),
        (Decimal("-10.00"), "BETWEEN BOUNDARIES"),
        (Decimal("0.00"), "MAXIMUM BOUNDARY"),
        (Decimal("20.00"), "ABOVE RANGE"),
    ]

    with SessionFactory.begin() as session:
        for index, (amount, description) in enumerate(transactions_by_amount):
            session.add(
                Transaction(
                    account_id=account_id,
                    import_fingerprint=f"{index + 40:064x}",
                    operation_date=date(2026, 8, 17),
                    value_date=date(2026, 8, 17),
                    amount=amount,
                    balance_after=Decimal("1000.00"),
                    bank_concept="TEST",
                    description=description,
                )
            )

    response = client.get(
        "/transactions",
        params={
            "amount_min": "-50.00",
            "amount_max": "0.00",
        },
    )

    assert response.status_code == 200
    assert {transaction["description"] for transaction in response.json()} == {
        "MINIMUM BOUNDARY",
        "BETWEEN BOUNDARIES",
        "MAXIMUM BOUNDARY",
    }


def test_list_transactions_rejects_reversed_amount_range(
    clean_database: None,
) -> None:
    response = client.get(
        "/transactions",
        params={
            "amount_min": "100.00",
            "amount_max": "-100.00",
        },
    )

    assert response.status_code == 422
    assert response.json() == {
        "detail": "amount_min must be less than or equal to amount_max"
    }


def test_import_ibercaja_xlsx_stores_transaction(clean_database: None) -> None:
    create_response = client.post(
        "/accounts",
        json={
            "name": "Main account",
            "bank_name": "Ibercaja",
            "currency": "EUR",
        },
    )

    assert create_response.status_code == 201

    account_id = create_response.json()["id"]
    workbook_contents = _create_fake_ibercaja_xlsx()
    import_response = client.post(
        f"/accounts/{account_id}/imports/ibercaja",
        files={
            "file": (
                "ibercaja.xlsx",
                workbook_contents,
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    )

    assert import_response.status_code == 201
    assert import_response.json() == {"imported": 1}

    list_response = client.get("/transactions")

    assert list_response.status_code == 200

    transactions = list_response.json()
    assert len(transactions) == 1
    assert transactions[0]["account_id"] == account_id
    assert transactions[0]["description"] == "MERCADONA CÑ DIEGO"
    assert transactions[0]["category"] == "food"


def test_import_ibercaja_xlsx_skips_duplicate_transaction(
    clean_database: None,
) -> None:
    create_response = client.post(
        "/accounts",
        json={
            "name": "Main account",
            "bank_name": "Ibercaja",
            "currency": "EUR",
        },
    )

    assert create_response.status_code == 201

    account_id = create_response.json()["id"]
    workbook_contents = _create_fake_ibercaja_xlsx()
    files = {
        "file": (
            "ibercaja.xlsx",
            workbook_contents,
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
    }

    first_response = client.post(
        f"/accounts/{account_id}/imports/ibercaja",
        files=files,
    )
    second_response = client.post(
        f"/accounts/{account_id}/imports/ibercaja",
        files=files,
    )

    assert first_response.status_code == 201
    assert first_response.json() == {"imported": 1}
    assert second_response.status_code == 201
    assert second_response.json() == {"imported": 0}

    list_response = client.get("/transactions")

    assert list_response.status_code == 200
    assert len(list_response.json()) == 1


def test_import_ibercaja_xlsx_rejects_unknown_account(clean_database: None) -> None:
    response = client.post(
        "/accounts/999999/imports/ibercaja",
        files={
            "file": (
                "ibercaja.xlsx",
                b"file contents are not parsed for a missing account",
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    )

    assert response.status_code == 404
    assert response.json() == {"detail": "Account not found"}


def test_import_ibercaja_xlsx_rejects_invalid_workbook(
    clean_database: None,
) -> None:
    create_response = client.post(
        "/accounts",
        json={
            "name": "Main account",
            "bank_name": "Ibercaja",
            "currency": "EUR",
        },
    )

    assert create_response.status_code == 201

    account_id = create_response.json()["id"]
    workbook = Workbook()
    sheet = workbook.active
    sheet["A1"] = "UNEXPECTED WORKBOOK"

    with BytesIO() as workbook_file:
        workbook.save(workbook_file)
        workbook.close()
        workbook_file.seek(0)

        response = client.post(
            f"/accounts/{account_id}/imports/ibercaja",
            files={
                "file": (
                    "ibercaja.xlsx",
                    workbook_file,
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                )
            },
        )

    assert response.status_code == 422
    assert response.json() == {"detail": "Invalid Ibercaja workbook"}


def test_import_ibercaja_xlsx_rejects_corrupt_file(clean_database: None) -> None:
    create_response = client.post(
        "/accounts",
        json={
            "name": "Main account",
            "bank_name": "Ibercaja",
            "currency": "EUR",
        },
    )

    assert create_response.status_code == 201

    account_id = create_response.json()["id"]
    response = client.post(
        f"/accounts/{account_id}/imports/ibercaja",
        files={
            "file": (
                "ibercaja.xlsx",
                b"this is plain text, not an XLSX workbook",
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    )

    assert response.status_code == 422
    assert response.json() == {"detail": "Invalid Ibercaja workbook"}


def test_import_revolut_investments_stores_and_lists_activities(
    clean_database: None,
) -> None:
    account_response = client.post(
        "/accounts",
        json={
            "name": "Investments",
            "bank_name": "Revolut",
            "currency": "EUR",
        },
    )
    account_id = account_response.json()["id"]

    import_response = client.post(
        f"/accounts/{account_id}/imports/revolut-investments",
        files={
            "file": (
                "revolut-investments.csv",
                _create_fake_revolut_investment_csv(),
                "text/csv",
            )
        },
    )

    assert import_response.status_code == 201
    assert import_response.json() == {"imported": 2}

    list_response = client.get(
        "/investment-activities",
        params={"account_id": account_id},
    )

    assert list_response.status_code == 200
    activities = list_response.json()
    assert len(activities) == 2
    assert activities[0]["activity_type"] == "DIVIDEND"
    assert activities[0]["ticker"] == "FAKE"
    assert Decimal(activities[0]["total_amount"]) == Decimal("0.25")
    assert activities[1]["activity_type"] == "BUY - MARKET"
    assert Decimal(activities[1]["quantity"]) == Decimal("0.5")
    assert Decimal(activities[1]["price_per_share"]) == Decimal(100)


def test_import_revolut_investments_skips_duplicates(clean_database: None) -> None:
    account_response = client.post(
        "/accounts",
        json={
            "name": "Investments",
            "bank_name": "Revolut",
            "currency": "EUR",
        },
    )
    account_id = account_response.json()["id"]
    files = {
        "file": (
            "revolut-investments.csv",
            _create_fake_revolut_investment_csv(),
            "text/csv",
        )
    }

    first_response = client.post(
        f"/accounts/{account_id}/imports/revolut-investments",
        files=files,
    )
    second_response = client.post(
        f"/accounts/{account_id}/imports/revolut-investments",
        files=files,
    )

    assert first_response.status_code == 201
    assert first_response.json() == {"imported": 2}
    assert second_response.status_code == 201
    assert second_response.json() == {"imported": 0}


def test_import_revolut_investments_rejects_invalid_csv(
    clean_database: None,
) -> None:
    account_response = client.post(
        "/accounts",
        json={
            "name": "Investments",
            "bank_name": "Revolut",
            "currency": "EUR",
        },
    )
    account_id = account_response.json()["id"]

    response = client.post(
        f"/accounts/{account_id}/imports/revolut-investments",
        files={"file": ("invalid.csv", b"Unexpected,Headers\n", "text/csv")},
    )

    assert response.status_code == 422
    assert response.json() == {"detail": "Invalid Revolut investment CSV"}


def test_import_revolut_investments_rejects_unknown_account(
    clean_database: None,
) -> None:
    response = client.post(
        "/accounts/999999/imports/revolut-investments",
        files={"file": ("investments.csv", b"not parsed", "text/csv")},
    )

    assert response.status_code == 404
    assert response.json() == {"detail": "Account not found"}


def test_investment_summary_uses_manual_price_for_profit_and_loss(
    clean_database: None,
) -> None:
    account_response = client.post(
        "/accounts",
        json={
            "name": "Investments",
            "bank_name": "Revolut",
            "currency": "EUR",
        },
    )
    account_id = account_response.json()["id"]
    import_response = client.post(
        f"/accounts/{account_id}/imports/revolut-investments",
        files={
            "file": (
                "revolut-investments.csv",
                _create_fake_revolut_investment_csv(),
                "text/csv",
            )
        },
    )

    assert import_response.status_code == 201

    unpriced_response = client.get(
        "/investment-summary",
        params={"account_id": account_id},
    )
    assert unpriced_response.status_code == 200
    assert unpriced_response.json()["positions"][0]["market_value"] is None
    assert unpriced_response.json()["currencies"][0]["total_result"] is None

    price_response = client.put(
        f"/accounts/{account_id}/investment-prices/FAKE",
        json={"price": "120"},
    )

    assert price_response.status_code == 200
    assert price_response.json()["ticker"] == "FAKE"
    assert Decimal(price_response.json()["price"]) == Decimal(120)

    summary_response = client.get(
        "/investment-summary",
        params={"account_id": account_id},
    )

    assert summary_response.status_code == 200
    position = summary_response.json()["positions"][0]
    assert Decimal(position["quantity"]) == Decimal("0.5")
    assert Decimal(position["remaining_cost"]) == Decimal(50)
    assert Decimal(position["market_value"]) == Decimal(60)
    assert Decimal(position["unrealized_pl"]) == Decimal(10)
    assert Decimal(position["realized_pl"]) == Decimal(0)
    assert Decimal(position["dividends"]) == Decimal("0.25")
    assert Decimal(position["total_result"]) == Decimal("10.25")

    currency_summary = summary_response.json()["currencies"][0]
    assert currency_summary["currency"] == "USD"
    assert Decimal(currency_summary["total_result"]) == Decimal("10.25")
    assert currency_summary["priced_positions"] == 1
    assert currency_summary["total_open_positions"] == 1


def test_update_investment_price_rejects_invalid_or_unknown_ticker(
    clean_database: None,
) -> None:
    account_response = client.post(
        "/accounts",
        json={
            "name": "Investments",
            "bank_name": "Revolut",
            "currency": "EUR",
        },
    )
    account_id = account_response.json()["id"]

    invalid_price_response = client.put(
        f"/accounts/{account_id}/investment-prices/FAKE",
        json={"price": "0"},
    )
    unknown_ticker_response = client.put(
        f"/accounts/{account_id}/investment-prices/UNKNOWN",
        json={"price": "100"},
    )

    assert invalid_price_response.status_code == 422
    assert unknown_ticker_response.status_code == 404
    assert unknown_ticker_response.json() == {"detail": "Ticker not found"}


def test_refresh_investment_prices_updates_usd_positions(
    clean_database: None,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    account_response = client.post(
        "/accounts",
        json={
            "name": "Investments",
            "bank_name": "Revolut",
            "currency": "EUR",
        },
    )
    account_id = account_response.json()["id"]
    client.post(
        f"/accounts/{account_id}/imports/revolut-investments",
        files={
            "file": (
                "revolut-investments.csv",
                _create_fake_revolut_investment_csv(),
                "text/csv",
            )
        },
    )
    captured_call: tuple[list[str], str] | None = None

    def fake_fetch(symbols: list[str], api_key: str) -> MarketPriceResult:
        nonlocal captured_call
        captured_call = (symbols, api_key)
        return MarketPriceResult(prices={"FAKE": Decimal(120)}, unavailable=[])

    monkeypatch.setenv("TWELVE_DATA_API_KEY", "fake-key")
    monkeypatch.setattr(investments_router, "fetch_twelve_data_prices", fake_fetch)

    response = client.post(f"/accounts/{account_id}/investment-prices/refresh")

    assert response.status_code == 200
    assert response.json() == {
        "updated": ["FAKE"],
        "unavailable": [],
        "manual_only": [],
    }
    assert captured_call == (["FAKE"], "fake-key")

    summary_response = client.get(
        "/investment-summary",
        params={"account_id": account_id},
    )
    assert Decimal(summary_response.json()["positions"][0]["current_price"]) == Decimal(
        120
    )


def test_refresh_investment_prices_requires_api_key(
    clean_database: None,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    account_response = client.post(
        "/accounts",
        json={
            "name": "Investments",
            "bank_name": "Revolut",
            "currency": "EUR",
        },
    )
    account_id = account_response.json()["id"]
    client.post(
        f"/accounts/{account_id}/imports/revolut-investments",
        files={
            "file": (
                "revolut-investments.csv",
                _create_fake_revolut_investment_csv(),
                "text/csv",
            )
        },
    )
    monkeypatch.delenv("TWELVE_DATA_API_KEY", raising=False)

    response = client.post(f"/accounts/{account_id}/investment-prices/refresh")

    assert response.status_code == 503
    assert response.json() == {"detail": "Market price API is not configured"}


def test_update_transaction_category_persists_category(
    clean_database: None,
) -> None:
    account_response = client.post(
        "/accounts",
        json={
            "name": "Main account",
            "bank_name": "Ibercaja",
            "currency": "EUR",
        },
    )
    assert account_response.status_code == 201
    account_id = account_response.json()["id"]

    with SessionFactory.begin() as session:
        transaction = Transaction(
            account_id=account_id,
            import_fingerprint="c" * 64,
            operation_date=date(2026, 8, 22),
            value_date=date(2026, 8, 22),
            amount=Decimal("-25.00"),
            balance_after=Decimal("1000.00"),
            bank_concept="CARD",
            description="FAKE RESTAURANT",
        )
        session.add(transaction)
        session.flush()
        transaction_id = transaction.id

    assert transaction_id is not None

    response = client.patch(
        f"/transactions/{transaction_id}/category",
        json={"category": "food"},
    )

    assert response.status_code == 200
    assert response.json()["category"] == "food"

    list_response = client.get("/transactions")

    assert list_response.status_code == 200
    assert list_response.json()[0]["category"] == "food"


def test_update_transaction_category_returns_404_when_transaction_does_not_exist(
    clean_database: None,
) -> None:
    response = client.patch(
        "/transactions/999999999/category",
        json={"category": "food"},
    )

    assert response.status_code == 404
    assert response.json() == {"detail": "Transaction not found"}


def test_update_transaction_category_rejects_empty_category(
    clean_database: None,
) -> None:
    response = client.patch(
        "/transactions/999999999/category",
        json={"category": ""},
    )

    assert response.status_code == 422
