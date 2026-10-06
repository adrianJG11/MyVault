from collections.abc import Iterator
from datetime import UTC, date, datetime, timedelta
from decimal import Decimal
from io import BytesIO
from xml.etree import ElementTree

import pytest
from alembic import command
from alembic.config import Config
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook
from sqlalchemy import delete, select
from sqlalchemy.exc import IntegrityError

from database import SessionFactory, engine
from investments import router as investments_router
from investments.market_prices import MarketPriceError, MarketPriceResult, MarketQuote
from main import app
from models import Account, InvestmentActivity, InvestmentPrice, Transaction

client = TestClient(app)


def _create_fake_ibercaja_xlsx(
    operation_date: str = "17-08-2026",
    balance: float = 14897.84,
) -> bytes:
    workbook = Workbook()
    sheet = workbook.active

    headers = [
        "Nº Orden",
        "Fecha Oper",
        "Fecha Valor",
        "Concepto",
        "Descripción",
        "Referencia",
        "Importe",
        "Saldo",
    ]
    for column, header in enumerate(headers, start=1):
        sheet.cell(row=6, column=column, value=header)

    fake_row = [
        1,
        operation_date,
        operation_date,
        "CARD",
        "MERCADONA CÑ DIEGO",
        "FAKE123",
        -42.64,
        balance,
    ]
    for column, value in enumerate(fake_row, start=1):
        sheet.cell(row=7, column=column, value=value)

    with BytesIO() as workbook_file:
        workbook.save(workbook_file)
        workbook.close()
        return workbook_file.getvalue()


def _create_fake_revolut_investment_csv() -> bytes:
    return (
        b"Date,Ticker,Type,Quantity,Price per share,Total Amount,Currency,FX Rate\n"
        b"2026-01-10T12:30:00Z,FAKE,BUY - MARKET,0.5,USD 100,USD 50,USD,1.2\n"
        b"2026-02-01T09:00:00Z,FAKE,DIVIDEND,,,USD 0.25,USD,1.1\n"
    )


def _create_fake_ibkr_xml(
    mark_price: str | None = None, report_date: str = "20260930"
) -> bytes:
    root = ElementTree.Element("FlexQueryResponse")
    statement = ElementTree.SubElement(root, "FlexStatement", accountId="TEST_ACCOUNT")
    trades = ElementTree.SubElement(statement, "Trades")
    common = {
        "accountId": "TEST_ACCOUNT",
        "assetCategory": "STK",
        "symbol": "FAKE",
        "currency": "USD",
        "fxRateToBase": "0.9",
    }
    ElementTree.SubElement(
        trades,
        "Trade",
        common
        | {
            "tradeID": "BUY123",
            "buySell": "BUY",
            "quantity": "2",
            "tradePrice": "10",
            "netCash": "-21",
            "dateTime": "20260110;123000",
        },
    )
    ElementTree.SubElement(
        trades,
        "Trade",
        common
        | {
            "tradeID": "SELL123",
            "buySell": "SELL",
            "quantity": "-1",
            "tradePrice": "12",
            "netCash": "11",
            "dateTime": "20260201;090000",
        },
    )
    if mark_price is not None:
        positions = ElementTree.SubElement(statement, "OpenPositions")
        ElementTree.SubElement(
            positions,
            "OpenPosition",
            common
            | {
                "levelOfDetail": "SUMMARY",
                "markPrice": mark_price,
                "reportDate": report_date,
            },
        )
    return ElementTree.tostring(root)


@pytest.fixture
def ibkr_account(clean_database: None) -> int:
    response = client.post(
        "/accounts",
        json={"name": "Investments", "bank_name": "IBKR", "currency": "EUR"},
    )
    assert response.status_code == 201
    return response.json()["id"]


@pytest.fixture
def clean_database() -> Iterator[None]:
    assert engine.url.database == "finanzas_test"

    with SessionFactory.begin() as session:
        session.execute(delete(InvestmentPrice))
        session.execute(delete(InvestmentActivity))
        session.execute(delete(Transaction))
        session.execute(delete(Account))

    yield

    with SessionFactory.begin() as session:
        session.execute(delete(InvestmentPrice))
        session.execute(delete(InvestmentActivity))
        session.execute(delete(Transaction))
        session.execute(delete(Account))


def test_health() -> None:
    response = client.get("/health")

    assert response.status_code == 200
    assert response.json() == {"status": "ok"}


def test_create_account_returns_account(clean_database: None) -> None:
    account_data = {
        "name": "Main account",
        "bank_name": "IGN",
        "currency": "EUR",
    }
    response = client.post("/accounts", json=account_data)
    response_data = response.json()

    assert response.status_code == 201
    assert isinstance(response_data["id"], int)
    assert response_data["id"] > 0
    assert response_data == {
        "id": response_data["id"],
        "current_balance": None,
        "balance_date": None,
        **account_data,
    }


def test_create_account_rejects_missing_currency() -> None:
    incomplete_account = {
        "name": "Main account",
        "bank_name": "ING",
    }

    response = client.post("/accounts", json=incomplete_account)

    assert response.status_code == 422


@pytest.mark.parametrize(
    ("field", "value"),
    [
        ("name", ""),
        ("name", "   "),
        ("name", "a" * 101),
        ("bank_name", ""),
        ("bank_name", "   "),
        ("bank_name", "a" * 101),
        ("currency", ""),
        ("currency", "EU"),
        ("currency", "EURO"),
        ("currency", "123"),
        ("currency", "€UR"),
        ("currency", None),
    ],
)
def test_create_account_rejects_invalid_fields(
    clean_database: None, field: str, value: str | None
) -> None:
    account_data = {"name": "Savings", "bank_name": "Example", "currency": "EUR"}
    account_data[field] = value

    response = client.post("/accounts", json=account_data)

    assert response.status_code == 422
    assert client.get("/accounts").json() == []


def test_create_account_normalizes_and_persists_fields(clean_database: None) -> None:
    response = client.post(
        "/accounts",
        json={"name": " Savings ", "bank_name": " Example ", "currency": " eur "},
    )

    assert response.status_code == 201
    account = response.json()
    assert account["name"] == "Savings"
    assert account["bank_name"] == "Example"
    assert account["currency"] == "EUR"
    assert client.get("/accounts").json() == [account]


def test_create_account_accepts_names_at_database_length_limit(
    clean_database: None,
) -> None:
    account_data = {"name": "a" * 100, "bank_name": "b" * 100, "currency": "USD"}

    response = client.post("/accounts", json=account_data)

    assert response.status_code == 201
    assert client.get("/accounts").json() == [response.json()]


def test_list_accounts_accepts_records_saved_before_input_validation(
    clean_database: None,
) -> None:
    with SessionFactory.begin() as session:
        session.add(Account(name="", bank_name="Example", currency="eur"))

    response = client.get("/accounts")

    assert response.status_code == 200
    assert response.json()[0]["name"] == ""
    assert response.json()[0]["currency"] == "eur"


def test_list_accounts_returns_created_account(clean_database: None) -> None:
    account_data = {
        "name": "Main account",
        "bank_name": "IGN",
        "currency": "EUR",
    }

    create_response = client.post("/accounts", json=account_data)

    assert create_response.status_code == 201

    created_account = create_response.json()
    list_response = client.get("/accounts")

    assert list_response.status_code == 200

    accounts = list_response.json()
    assert accounts == [created_account]


def test_list_empty_database(clean_database: None) -> None:
    list_response = client.get("/accounts")
    assert list_response.status_code == 200
    accounts = list_response.json()
    assert accounts == []


def test_list_transactions_returns_empty_list(clean_database: None) -> None:
    list_response = client.get("/transactions")
    assert list_response.status_code == 200
    transactions = list_response.json()
    assert transactions == []


def test_list_transactions_returns_stored_transaction(clean_database: None) -> None:
    account_data = {
        "name": "Main account",
        "bank_name": "ING",
        "currency": "EUR",
    }

    create_response = client.post("/accounts", json=account_data)

    assert create_response.status_code == 201

    account_id = create_response.json()["id"]

    with SessionFactory.begin() as session:
        transaction = Transaction(
            account_id=account_id,
            import_fingerprint="a" * 64,
            operation_date=date(2026, 8, 17),
            value_date=date(2026, 8, 17),
            amount=Decimal("-42.64"),
            balance_after=Decimal("14897.84"),
            bank_concept="CARD",
            description="FAKE SUPERMARKET",
        )
        session.add(transaction)

    list_response = client.get("/transactions")

    assert list_response.status_code == 200

    transactions = list_response.json()
    assert len(transactions) == 1

    transaction_data = transactions[0]
    assert isinstance(transaction_data["id"], int)
    assert transaction_data == {
        "id": transaction_data["id"],
        "account_id": account_id,
        "operation_date": "2026-08-17",
        "value_date": "2026-08-17",
        "amount": "-42.64",
        "balance_after": "14897.84",
        "bank_concept": "CARD",
        "description": "FAKE SUPERMARKET",
        "category": None,
    }


def test_list_transactions_filters_by_account_id(clean_database: None) -> None:
    first_account_response = client.post(
        "/accounts",
        json={
            "name": "First account",
            "bank_name": "Ibercaja",
            "currency": "EUR",
        },
    )
    second_account_response = client.post(
        "/accounts",
        json={
            "name": "Second account",
            "bank_name": "Other bank",
            "currency": "EUR",
        },
    )

    assert first_account_response.status_code == 201
    assert second_account_response.status_code == 201

    first_account_id = first_account_response.json()["id"]
    second_account_id = second_account_response.json()["id"]

    with SessionFactory.begin() as session:
        session.add_all(
            [
                Transaction(
                    account_id=first_account_id,
                    import_fingerprint="b" * 64,
                    operation_date=date(2026, 8, 17),
                    value_date=date(2026, 8, 17),
                    amount=Decimal("-42.64"),
                    balance_after=Decimal("14897.84"),
                    bank_concept="CARD",
                    description="FIRST ACCOUNT PURCHASE",
                ),
                Transaction(
                    account_id=second_account_id,
                    import_fingerprint="c" * 64,
                    operation_date=date(2026, 8, 17),
                    value_date=date(2026, 8, 17),
                    amount=Decimal("-18.57"),
                    balance_after=Decimal("1000.00"),
                    bank_concept="CARD",
                    description="SECOND ACCOUNT PURCHASE",
                ),
            ]
        )

    response = client.get(
        "/transactions",
        params={"account_id": first_account_id},
    )

    assert response.status_code == 200

    transactions = response.json()
    assert len(transactions) == 1
    assert transactions[0]["account_id"] == first_account_id
    assert transactions[0]["description"] == "FIRST ACCOUNT PURCHASE"


def test_list_transactions_filters_by_operation_date_range(
    clean_database: None,
) -> None:
    create_response = client.post(
        "/accounts",
        json={
            "name": "Main account",
            "bank_name": "Ibercaja",
            "currency": "EUR",
        },
    )

    assert create_response.status_code == 201

    account_id = create_response.json()["id"]
    dated_transactions = [
        (date(2026, 8, 9), "BEFORE RANGE"),
        (date(2026, 8, 10), "START BOUNDARY"),
        (date(2026, 8, 20), "END BOUNDARY"),
        (date(2026, 8, 21), "AFTER RANGE"),
    ]

    with SessionFactory.begin() as session:
        for index, (operation_date, description) in enumerate(dated_transactions):
            session.add(
                Transaction(
                    account_id=account_id,
                    import_fingerprint=f"{index:064x}",
                    operation_date=operation_date,
                    value_date=operation_date,
                    amount=Decimal("-1.00"),
                    balance_after=Decimal("1000.00") - index,
                    bank_concept="TEST",
                    description=description,
                )
            )

    response = client.get(
        "/transactions",
        params={
            "date_from": "2026-08-10",
            "date_to": "2026-08-20",
        },
    )

    assert response.status_code == 200

    transactions = response.json()
    assert len(transactions) == 2
    assert {transaction["description"] for transaction in transactions} == {
        "START BOUNDARY",
        "END BOUNDARY",
    }


def test_list_transactions_rejects_reversed_date_range(
    clean_database: None,
) -> None:
    response = client.get(
        "/transactions",
        params={
            "date_from": "2026-08-20",
            "date_to": "2026-08-10",
        },
    )

    assert response.status_code == 422
    assert response.json() == {"detail": "date_from must be before or equal to date_to"}


def test_list_transactions_orders_newest_first(clean_database: None) -> None:
    create_response = client.post(
        "/accounts",
        json={
            "name": "Main account",
            "bank_name": "Ibercaja",
            "currency": "EUR",
        },
    )

    assert create_response.status_code == 201

    account_id = create_response.json()["id"]
    dated_transactions = [
        (date(2026, 8, 10), "OLDEST"),
        (date(2026, 8, 20), "NEWEST FIRST"),
        (date(2026, 8, 20), "NEWEST SECOND"),
    ]

    with SessionFactory.begin() as session:
        for index, (operation_date, description) in enumerate(dated_transactions):
            session.add(
                Transaction(
                    account_id=account_id,
                    import_fingerprint=f"{index + 20:064x}",
                    operation_date=operation_date,
                    value_date=operation_date,
                    amount=Decimal("-1.00"),
                    balance_after=Decimal("1000.00") - index,
                    bank_concept="TEST",
                    description=description,
                )
            )

    response = client.get("/transactions")

    assert response.status_code == 200
    assert [transaction["description"] for transaction in response.json()] == [
        "NEWEST SECOND",
        "NEWEST FIRST",
        "OLDEST",
    ]


def test_list_transactions_filters_by_description(clean_database: None) -> None:
    create_response = client.post(
        "/accounts",
        json={
            "name": "Main account",
            "bank_name": "Ibercaja",
            "currency": "EUR",
        },
    )

    assert create_response.status_code == 201

    account_id = create_response.json()["id"]

    with SessionFactory.begin() as session:
        session.add_all(
            [
                Transaction(
                    account_id=account_id,
                    import_fingerprint=f"{30:064x}",
                    operation_date=date(2026, 8, 17),
                    value_date=date(2026, 8, 17),
                    amount=Decimal("-42.64"),
                    balance_after=Decimal("14897.84"),
                    bank_concept="CARD",
                    description="MERCADONA CÑ DIEGO",
                ),
                Transaction(
                    account_id=account_id,
                    import_fingerprint=f"{31:064x}",
                    operation_date=date(2026, 8, 16),
                    value_date=date(2026, 8, 16),
                    amount=Decimal("2000.00"),
                    balance_after=Decimal("14940.48"),
                    bank_concept="TRANSFER",
                    description="SALARY",
                ),
            ]
        )

    response = client.get(
        "/transactions",
        params={"description": "mercadona"},
    )

    assert response.status_code == 200

    transactions = response.json()
    assert len(transactions) == 1
    assert transactions[0]["description"] == "MERCADONA CÑ DIEGO"


def test_list_transactions_filters_by_category(clean_database: None) -> None:
    create_response = client.post(
        "/accounts",
        json={
            "name": "Main account",
            "bank_name": "Ibercaja",
            "currency": "EUR",
        },
    )

    assert create_response.status_code == 201

    account_id = create_response.json()["id"]

    with SessionFactory.begin() as session:
        session.add_all(
            [
                Transaction(
                    account_id=account_id,
                    import_fingerprint=f"{90:064x}",
                    operation_date=date(2026, 8, 17),
                    value_date=date(2026, 8, 17),
                    amount=Decimal("-686.11"),
                    balance_after=Decimal("14776.84"),
                    bank_concept="TRANSFER",
                    description="RENT PAYMENT",
                    category="housing",
                ),
                Transaction(
                    account_id=account_id,
                    import_fingerprint=f"{91:064x}",
                    operation_date=date(2026, 8, 16),
                    value_date=date(2026, 8, 16),
                    amount=Decimal("-42.64"),
                    balance_after=Decimal("15462.95"),
                    bank_concept="CARD",
                    description="SUPERMARKET PURCHASE",
                    category="food",
                ),
            ]
        )

    response = client.get(
        "/transactions",
        params={"category": "housing"},
    )

    assert response.status_code == 200

    transactions = response.json()
    assert len(transactions) == 1
    assert transactions[0]["description"] == "RENT PAYMENT"
    assert transactions[0]["category"] == "housing"


def test_list_transactions_filters_by_uncategorized(clean_database: None) -> None:
    create_response = client.post(
        "/accounts",
        json={
            "name": "Main account",
            "bank_name": "Ibercaja",
            "currency": "EUR",
        },
    )

    assert create_response.status_code == 201

    account_id = create_response.json()["id"]

    with SessionFactory.begin() as session:
        session.add_all(
            [
                Transaction(
                    account_id=account_id,
                    import_fingerprint=f"{92:064x}",
                    operation_date=date(2026, 8, 17),
                    value_date=date(2026, 8, 17),
                    amount=Decimal("-15.00"),
                    balance_after=Decimal("1000.00"),
                    bank_concept="CARD",
                    description="UNKNOWN PURCHASE",
                    category=None,
                ),
                Transaction(
                    account_id=account_id,
                    import_fingerprint=f"{93:064x}",
                    operation_date=date(2026, 8, 16),
                    value_date=date(2026, 8, 16),
                    amount=Decimal("-42.64"),
                    balance_after=Decimal("1015.00"),
                    bank_concept="CARD",
                    description="SUPERMARKET PURCHASE",
                    category="food",
                ),
            ]
        )

    response = client.get(
        "/transactions",
        params={"category": "uncategorized"},
    )

    assert response.status_code == 200

    transactions = response.json()
    assert len(transactions) == 1
    assert transactions[0]["description"] == "UNKNOWN PURCHASE"
    assert transactions[0]["category"] is None


def test_list_transactions_filters_by_amount_range(clean_database: None) -> None:
    create_response = client.post(
        "/accounts",
        json={
            "name": "Main account",
            "bank_name": "Ibercaja",
            "currency": "EUR",
        },
    )

    assert create_response.status_code == 201

    account_id = create_response.json()["id"]
    transactions_by_amount = [
        (Decimal("-100.00"), "BELOW RANGE"),
        (Decimal("-50.00"), "MINIMUM BOUNDARY"),
        (Decimal("-10.00"), "BETWEEN BOUNDARIES"),
        (Decimal("0.00"), "MAXIMUM BOUNDARY"),
        (Decimal("20.00"), "ABOVE RANGE"),
    ]

    with SessionFactory.begin() as session:
        for index, (amount, description) in enumerate(transactions_by_amount):
            session.add(
                Transaction(
                    account_id=account_id,
                    import_fingerprint=f"{index + 40:064x}",
                    operation_date=date(2026, 8, 17),
                    value_date=date(2026, 8, 17),
                    amount=amount,
                    balance_after=Decimal("1000.00"),
                    bank_concept="TEST",
                    description=description,
                )
            )

    response = client.get(
        "/transactions",
        params={
            "amount_min": "-50.00",
            "amount_max": "0.00",
        },
    )

    assert response.status_code == 200
    assert {transaction["description"] for transaction in response.json()} == {
        "MINIMUM BOUNDARY",
        "BETWEEN BOUNDARIES",
        "MAXIMUM BOUNDARY",
    }


def test_list_transactions_rejects_reversed_amount_range(
    clean_database: None,
) -> None:
    response = client.get(
        "/transactions",
        params={
            "amount_min": "100.00",
            "amount_max": "-100.00",
        },
    )

    assert response.status_code == 422
    assert response.json() == {
        "detail": "amount_min must be less than or equal to amount_max"
    }


def test_import_ibercaja_xlsx_stores_transaction_and_updates_balance(
    clean_database: None,
) -> None:
    create_response = client.post(
        "/accounts",
        json={
            "name": "Main account",
            "bank_name": "Ibercaja",
            "currency": "EUR",
        },
    )

    assert create_response.status_code == 201

    account_id = create_response.json()["id"]
    workbook_contents = _create_fake_ibercaja_xlsx()
    import_response = client.post(
        f"/accounts/{account_id}/imports/ibercaja",
        files={
            "file": (
                "ibercaja.xlsx",
                workbook_contents,
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    )

    assert import_response.status_code == 201
    assert import_response.json() == {"imported": 1, "skipped": 0}

    list_response = client.get("/transactions")

    assert list_response.status_code == 200

    transactions = list_response.json()
    assert len(transactions) == 1
    assert transactions[0]["account_id"] == account_id
    assert transactions[0]["description"] == "MERCADONA CÑ DIEGO"
    assert transactions[0]["category"] == "food"

    accounts_response = client.get("/accounts")

    assert accounts_response.status_code == 200

    account = accounts_response.json()[0]
    assert account["current_balance"] == "14897.84"
    assert account["balance_date"] == "2026-08-17"


def test_import_older_transaction_does_not_replace_latest_balance(
    clean_database: None,
) -> None:
    create_response = client.post(
        "/accounts",
        json={
            "name": "Main account",
            "bank_name": "Ibercaja",
            "currency": "EUR",
        },
    )

    assert create_response.status_code == 201

    account_id = create_response.json()["id"]
    newer_workbook = _create_fake_ibercaja_xlsx("17-08-2026", 1000.00)
    older_workbook = _create_fake_ibercaja_xlsx("10-08-2026", 900.00)

    for workbook_contents in (newer_workbook, older_workbook):
        import_response = client.post(
            f"/accounts/{account_id}/imports/ibercaja",
            files={
                "file": (
                    "ibercaja.xlsx",
                    workbook_contents,
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                )
            },
        )

        assert import_response.status_code == 201
        assert import_response.json() == {"imported": 1, "skipped": 0}

    accounts_response = client.get("/accounts")

    assert accounts_response.status_code == 200

    account = accounts_response.json()[0]
    assert account["current_balance"] == "1000.00"
    assert account["balance_date"] == "2026-08-17"


@pytest.mark.parametrize("duplicate_in_file", [False, True])
def test_import_ibercaja_xlsx_skips_duplicate_transaction(
    clean_database: None,
    duplicate_in_file: bool,
) -> None:
    create_response = client.post(
        "/accounts",
        json={
            "name": "Main account",
            "bank_name": "Ibercaja",
            "currency": "EUR",
        },
    )

    assert create_response.status_code == 201

    account_id = create_response.json()["id"]
    workbook_contents = _create_fake_ibercaja_xlsx()
    if duplicate_in_file:
        workbook = load_workbook(BytesIO(workbook_contents))
        sheet = workbook.active
        sheet.append([cell.value for cell in sheet[7]])
        with BytesIO() as workbook_file:
            workbook.save(workbook_file)
            workbook_contents = workbook_file.getvalue()
        workbook.close()

    files = {
        "file": (
            "ibercaja.xlsx",
            workbook_contents,
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
    }

    first_response = client.post(
        f"/accounts/{account_id}/imports/ibercaja",
        files=files,
    )
    second_response = client.post(
        f"/accounts/{account_id}/imports/ibercaja",
        files=files,
    )

    assert first_response.status_code == 201
    assert first_response.json() == {"imported": 1, "skipped": int(duplicate_in_file)}
    assert second_response.status_code == 201
    assert second_response.json() == {
        "imported": 0,
        "skipped": 2 if duplicate_in_file else 1,
    }

    list_response = client.get("/transactions")

    assert list_response.status_code == 200
    assert len(list_response.json()) == 1


def test_import_ibercaja_xlsx_rejects_unknown_account(clean_database: None) -> None:
    response = client.post(
        "/accounts/999999/imports/ibercaja",
        files={
            "file": (
                "ibercaja.xlsx",
                b"file contents are not parsed for a missing account",
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    )

    assert response.status_code == 404
    assert response.json() == {"detail": "Account not found"}


def test_import_ibercaja_xlsx_rejects_invalid_workbook(
    clean_database: None,
) -> None:
    create_response = client.post(
        "/accounts",
        json={
            "name": "Main account",
            "bank_name": "Ibercaja",
            "currency": "EUR",
        },
    )

    assert create_response.status_code == 201

    account_id = create_response.json()["id"]
    workbook = Workbook()
    sheet = workbook.active
    sheet["A1"] = "UNEXPECTED WORKBOOK"

    with BytesIO() as workbook_file:
        workbook.save(workbook_file)
        workbook.close()
        workbook_file.seek(0)

        response = client.post(
            f"/accounts/{account_id}/imports/ibercaja",
            files={
                "file": (
                    "ibercaja.xlsx",
                    workbook_file,
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                )
            },
        )

    assert response.status_code == 422
    assert response.json() == {"detail": "Invalid Ibercaja workbook"}


def test_import_ibercaja_xlsx_rejects_corrupt_file(clean_database: None) -> None:
    create_response = client.post(
        "/accounts",
        json={
            "name": "Main account",
            "bank_name": "Ibercaja",
            "currency": "EUR",
        },
    )

    assert create_response.status_code == 201

    account_id = create_response.json()["id"]
    response = client.post(
        f"/accounts/{account_id}/imports/ibercaja",
        files={
            "file": (
                "ibercaja.xlsx",
                b"this is plain text, not an XLSX workbook",
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    )

    assert response.status_code == 422
    assert response.json() == {"detail": "Invalid Ibercaja workbook"}


def test_import_revolut_investments_stores_and_lists_activities(
    clean_database: None,
) -> None:
    account_response = client.post(
        "/accounts",
        json={
            "name": "Investments",
            "bank_name": "Revolut",
            "currency": "EUR",
        },
    )
    account_id = account_response.json()["id"]

    import_response = client.post(
        f"/accounts/{account_id}/imports/revolut-investments",
        files={
            "file": (
                "revolut-investments.csv",
                _create_fake_revolut_investment_csv(),
                "text/csv",
            )
        },
    )

    assert import_response.status_code == 201
    assert import_response.json() == {"imported": 2}

    list_response = client.get(
        "/investment-activities",
        params={"account_id": account_id},
    )

    assert list_response.status_code == 200
    activities = list_response.json()
    assert len(activities) == 2
    assert activities[0]["activity_type"] == "DIVIDEND"
    assert activities[0]["ticker"] == "FAKE"
    assert Decimal(activities[0]["total_amount"]) == Decimal("0.25")
    assert activities[1]["activity_type"] == "BUY - MARKET"
    assert Decimal(activities[1]["quantity"]) == Decimal("0.5")
    assert Decimal(activities[1]["price_per_share"]) == Decimal(100)


def test_import_revolut_investments_skips_duplicates(clean_database: None) -> None:
    account_response = client.post(
        "/accounts",
        json={
            "name": "Investments",
            "bank_name": "Revolut",
            "currency": "EUR",
        },
    )
    account_id = account_response.json()["id"]
    files = {
        "file": (
            "revolut-investments.csv",
            _create_fake_revolut_investment_csv(),
            "text/csv",
        )
    }

    first_response = client.post(
        f"/accounts/{account_id}/imports/revolut-investments",
        files=files,
    )
    second_response = client.post(
        f"/accounts/{account_id}/imports/revolut-investments",
        files=files,
    )

    assert first_response.status_code == 201
    assert first_response.json() == {"imported": 2}
    assert second_response.status_code == 201
    assert second_response.json() == {"imported": 0}


def test_import_revolut_investments_rejects_invalid_csv(
    clean_database: None,
) -> None:
    account_response = client.post(
        "/accounts",
        json={
            "name": "Investments",
            "bank_name": "Revolut",
            "currency": "EUR",
        },
    )
    account_id = account_response.json()["id"]

    response = client.post(
        f"/accounts/{account_id}/imports/revolut-investments",
        files={"file": ("invalid.csv", b"Unexpected,Headers\n", "text/csv")},
    )

    assert response.status_code == 422
    assert response.json() == {"detail": "Invalid Revolut investment CSV"}


def test_import_revolut_investments_rejects_unknown_account(
    clean_database: None,
) -> None:
    response = client.post(
        "/accounts/999999/imports/revolut-investments",
        files={"file": ("investments.csv", b"not parsed", "text/csv")},
    )

    assert response.status_code == 404
    assert response.json() == {"detail": "Account not found"}


def test_import_ibkr_stores_lists_and_calculates_activities(ibkr_account: int) -> None:
    response = client.post(
        f"/accounts/{ibkr_account}/imports/ibkr-investments",
        files={"file": ("report.xml", _create_fake_ibkr_xml(), "application/xml")},
    )
    assert response.status_code == 201
    assert response.json() == {"imported": 2, "prices_updated": 0}

    activities = client.get(
        "/investment-activities", params={"account_id": ibkr_account}
    ).json()
    assert len(activities) == 2
    assert activities[0]["activity_type"] == "SELL - MARKET"
    assert Decimal(activities[0]["quantity"]) == Decimal(1)
    assert Decimal(activities[0]["total_amount"]) == Decimal(11)
    assert Decimal(activities[0]["fx_rate"]) == Decimal("0.9")
    assert activities[1]["occurred_at"] == "2026-01-10T17:30:00Z"
    assert Decimal(activities[1]["total_amount"]) == Decimal(21)

    summary = client.get("/investment-summary", params={"account_id": ibkr_account})
    assert summary.status_code == 200
    position = summary.json()["positions"][0]
    assert Decimal(position["quantity"]) == Decimal(1)
    assert Decimal(position["remaining_cost"]) == Decimal("10.5")
    assert Decimal(position["realized_pl"]) == Decimal("0.5")


def test_import_ibkr_skips_duplicates_in_file_and_repeat_uploads(
    ibkr_account: int,
) -> None:
    root = ElementTree.fromstring(_create_fake_ibkr_xml())
    trades = root.find(".//Trades")
    assert trades is not None
    trades.extend(list(trades))
    files = {"file": ("report.xml", ElementTree.tostring(root), "application/xml")}

    first = client.post(
        f"/accounts/{ibkr_account}/imports/ibkr-investments", files=files
    )
    second = client.post(
        f"/accounts/{ibkr_account}/imports/ibkr-investments", files=files
    )

    assert first.status_code == second.status_code == 201
    assert first.json() == {"imported": 2, "prices_updated": 0}
    assert second.json() == {"imported": 0, "prices_updated": 0}
    assert (
        len(
            client.get(
                "/investment-activities", params={"account_id": ibkr_account}
            ).json()
        )
        == 2
    )


def test_import_ibkr_does_not_merge_distinct_identical_executions(
    ibkr_account: int,
) -> None:
    root = ElementTree.fromstring(_create_fake_ibkr_xml())
    trades = root.find(".//Trades")
    assert trades is not None
    duplicate = ElementTree.fromstring(ElementTree.tostring(trades[0]))
    duplicate.attrib["tradeID"] = "BUY124"
    trades.append(duplicate)

    response = client.post(
        f"/accounts/{ibkr_account}/imports/ibkr-investments",
        files={"file": ("report.xml", ElementTree.tostring(root), "application/xml")},
    )
    assert response.status_code == 201
    assert response.json() == {"imported": 3, "prices_updated": 0}


def test_import_ibkr_rejects_invalid_trade_without_partial_import(
    ibkr_account: int,
) -> None:
    root = ElementTree.fromstring(_create_fake_ibkr_xml())
    list(root.iter("Trade"))[1].attrib["quantity"] = "abc"

    response = client.post(
        f"/accounts/{ibkr_account}/imports/ibkr-investments",
        files={"file": ("report.xml", ElementTree.tostring(root), "application/xml")},
    )
    assert response.status_code == 422
    assert response.json() == {"detail": "Invalid IBKR report"}
    assert (
        client.get("/investment-activities", params={"account_id": ibkr_account}).json()
        == []
    )


def test_import_ibkr_rolls_back_earlier_batches_on_numeric_overflow(
    ibkr_account: int,
) -> None:
    original = next(ElementTree.fromstring(_create_fake_ibkr_xml()).iter("Trade"))
    root = ElementTree.Element("FlexQueryResponse")
    for index in range(501):
        attributes = original.attrib | {"tradeID": f"BUY{index}"}
        if index == 500:
            attributes["quantity"] = "1e40"
        ElementTree.SubElement(root, "Trade", attributes)

    response = client.post(
        f"/accounts/{ibkr_account}/imports/ibkr-investments",
        files={"file": ("report.xml", ElementTree.tostring(root), "application/xml")},
    )
    assert response.status_code == 422
    assert response.json() == {"detail": "Investment values cannot be stored"}
    assert (
        client.get("/investment-activities", params={"account_id": ibkr_account}).json()
        == []
    )


@pytest.mark.parametrize("report", [b"not XML", b"<WrongRoot/>"])
def test_import_ibkr_rejects_invalid_xml(ibkr_account: int, report: bytes) -> None:
    response = client.post(
        f"/accounts/{ibkr_account}/imports/ibkr-investments",
        files={"file": ("report.xml", report, "application/xml")},
    )
    assert response.status_code == 422
    assert response.json() == {"detail": "Invalid IBKR report"}


def test_import_ibkr_accepts_empty_report(ibkr_account: int) -> None:
    response = client.post(
        f"/accounts/{ibkr_account}/imports/ibkr-investments",
        files={"file": ("report.xml", b"<FlexQueryResponse/>", "application/xml")},
    )
    assert response.status_code == 201
    assert response.json() == {"imported": 0, "prices_updated": 0}


def test_import_ibkr_rejects_unknown_account(clean_database: None) -> None:
    response = client.post(
        "/accounts/999999/imports/ibkr-investments",
        files={"file": ("report.xml", _create_fake_ibkr_xml(), "application/xml")},
    )
    assert response.status_code == 404
    assert response.json() == {"detail": "Account not found"}


def test_import_ibkr_prices_calculates_value_and_exposes_report_date(
    ibkr_account: int,
) -> None:
    response = client.post(
        f"/accounts/{ibkr_account}/imports/ibkr-investments",
        files={"file": ("report.xml", _create_fake_ibkr_xml("15"), "application/xml")},
    )
    assert response.status_code == 201
    assert response.json() == {"imported": 2, "prices_updated": 1}

    summary = client.get("/investment-summary", params={"account_id": ibkr_account})
    assert summary.status_code == 200
    position = summary.json()["positions"][0]
    assert Decimal(position["market_value"]) == Decimal(15)
    assert Decimal(position["unrealized_pl"]) == Decimal("4.5")
    assert position["price_source"] == "ibkr"
    assert position["price_as_of"] == "2026-09-30"
    assert position["price_updated_at"] is not None


def test_ibkr_prices_support_sxrv_and_vwce_without_external_quotes(
    ibkr_account: int,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    root = ElementTree.fromstring(_create_fake_ibkr_xml("15"))
    for node in root.iter():
        if "symbol" in node.attrib:
            node.attrib.update(symbol="SXRV", currency="EUR")
    trades = root.find(".//Trades")
    positions = root.find(".//OpenPositions")
    assert trades is not None and positions is not None
    ElementTree.SubElement(
        trades, "Trade", trades[0].attrib | {"tradeID": "VWCE_BUY", "symbol": "VWCE"}
    )
    ElementTree.SubElement(
        positions,
        "OpenPosition",
        positions[0].attrib | {"symbol": "VWCE", "markPrice": "20"},
    )

    response = client.post(
        f"/accounts/{ibkr_account}/imports/ibkr-investments",
        files={"file": ("report.xml", ElementTree.tostring(root), "application/xml")},
    )
    assert response.status_code == 201
    assert response.json() == {"imported": 3, "prices_updated": 2}
    summary = client.get(
        "/investment-summary", params={"account_id": ibkr_account}
    ).json()
    assert {position["ticker"] for position in summary["positions"]} == {"SXRV", "VWCE"}
    assert Decimal(summary["currencies"][0]["market_value"]) == Decimal(55)

    def unavailable_quotes(
        symbols: dict[str, str], currencies: dict[str, str]
    ) -> MarketPriceResult:
        assert symbols == {"SXRV": "SXRV.DE", "VWCE": "VWCE.DE"}
        assert currencies == {"SXRV": "EUR", "VWCE": "EUR"}
        return MarketPriceResult(prices={}, unavailable=["SXRV", "VWCE"])

    monkeypatch.setattr(investments_router, "fetch_yahoo_prices", unavailable_quotes)
    refresh = client.post(f"/accounts/{ibkr_account}/investment-prices/refresh")
    assert refresh.status_code == 200
    assert refresh.json()["unavailable"] == ["SXRV", "VWCE"]
    assert refresh.json()["manual_only"] == []
    assert (
        client.get("/investment-summary", params={"account_id": ibkr_account}).json()
        == summary
    )


def test_import_ibkr_updates_prices_without_new_trades_and_ignores_older_reports(
    ibkr_account: int,
) -> None:
    for mark_price, report_date, expected in [
        ("15", "20260929", {"imported": 2, "prices_updated": 1}),
        ("16", "20260930", {"imported": 0, "prices_updated": 1}),
        ("16", "20260930", {"imported": 0, "prices_updated": 0}),
        ("14", "20260928", {"imported": 0, "prices_updated": 0}),
    ]:
        response = client.post(
            f"/accounts/{ibkr_account}/imports/ibkr-investments",
            files={
                "file": (
                    "report.xml",
                    _create_fake_ibkr_xml(mark_price, report_date),
                    "application/xml",
                )
            },
        )
        assert response.status_code == 201
        assert response.json() == expected

    position = client.get(
        "/investment-summary", params={"account_id": ibkr_account}
    ).json()["positions"][0]
    assert Decimal(position["current_price"]) == Decimal(16)
    assert position["price_as_of"] == "2026-09-30"


def test_import_ibkr_position_only_report_refreshes_existing_holdings(
    ibkr_account: int,
) -> None:
    client.post(
        f"/accounts/{ibkr_account}/imports/ibkr-investments",
        files={"file": ("report.xml", _create_fake_ibkr_xml(), "application/xml")},
    )
    root = ElementTree.fromstring(_create_fake_ibkr_xml("15"))
    statement = root.find("FlexStatement")
    assert statement is not None
    trades = statement.find("Trades")
    assert trades is not None
    statement.remove(trades)
    response = client.post(
        f"/accounts/{ibkr_account}/imports/ibkr-investments",
        files={"file": ("prices.xml", ElementTree.tostring(root), "application/xml")},
    )
    assert response.status_code == 201
    assert response.json() == {"imported": 0, "prices_updated": 1}


def test_old_ibkr_report_does_not_replace_recent_manual_price(
    ibkr_account: int,
) -> None:
    client.post(
        f"/accounts/{ibkr_account}/imports/ibkr-investments",
        files={"file": ("report.xml", _create_fake_ibkr_xml(), "application/xml")},
    )
    client.put(f"/accounts/{ibkr_account}/investment-prices/FAKE", json={"price": "20"})
    response = client.post(
        f"/accounts/{ibkr_account}/imports/ibkr-investments",
        files={"file": ("report.xml", _create_fake_ibkr_xml("15"), "application/xml")},
    )
    assert response.status_code == 201
    assert response.json() == {"imported": 0, "prices_updated": 0}
    position = client.get(
        "/investment-summary", params={"account_id": ibkr_account}
    ).json()["positions"][0]
    assert Decimal(position["current_price"]) == Decimal(20)
    assert position["price_source"] == "manual"
    assert position["price_as_of"] is None


@pytest.mark.parametrize("mark_price", ["NaN", "1e40"])
def test_invalid_ibkr_price_does_not_save_trades_or_prices(
    ibkr_account: int,
    mark_price: str,
) -> None:
    response = client.post(
        f"/accounts/{ibkr_account}/imports/ibkr-investments",
        files={
            "file": ("report.xml", _create_fake_ibkr_xml(mark_price), "application/xml")
        },
    )
    assert response.status_code == 422
    assert (
        client.get("/investment-activities", params={"account_id": ibkr_account}).json()
        == []
    )
    with SessionFactory() as session:
        assert session.query(InvestmentPrice).count() == 0


@pytest.mark.parametrize("field,value", [("currency", "EUR"), ("symbol", "UNKNOWN")])
def test_ibkr_price_requires_matching_trade_history(
    ibkr_account: int,
    field: str,
    value: str,
) -> None:
    root = ElementTree.fromstring(_create_fake_ibkr_xml("15"))
    next(root.iter("OpenPosition")).attrib[field] = value
    response = client.post(
        f"/accounts/{ibkr_account}/imports/ibkr-investments",
        files={"file": ("report.xml", ElementTree.tostring(root), "application/xml")},
    )
    assert response.status_code == 422
    assert (
        client.get("/investment-activities", params={"account_id": ibkr_account}).json()
        == []
    )


def test_investment_summary_uses_manual_price_for_profit_and_loss(
    clean_database: None,
) -> None:
    account_response = client.post(
        "/accounts",
        json={
            "name": "Investments",
            "bank_name": "Revolut",
            "currency": "EUR",
        },
    )
    account_id = account_response.json()["id"]
    import_response = client.post(
        f"/accounts/{account_id}/imports/revolut-investments",
        files={
            "file": (
                "revolut-investments.csv",
                _create_fake_revolut_investment_csv(),
                "text/csv",
            )
        },
    )

    assert import_response.status_code == 201

    unpriced_response = client.get(
        "/investment-summary",
        params={"account_id": account_id},
    )
    assert unpriced_response.status_code == 200
    assert unpriced_response.json()["positions"][0]["market_value"] is None
    assert unpriced_response.json()["currencies"][0]["total_result"] is None

    price_response = client.put(
        f"/accounts/{account_id}/investment-prices/FAKE",
        json={"price": "120"},
    )

    assert price_response.status_code == 200
    assert price_response.json()["ticker"] == "FAKE"
    assert Decimal(price_response.json()["price"]) == Decimal(120)

    summary_response = client.get(
        "/investment-summary",
        params={"account_id": account_id},
    )

    assert summary_response.status_code == 200
    position = summary_response.json()["positions"][0]
    assert Decimal(position["quantity"]) == Decimal("0.5")
    assert Decimal(position["remaining_cost"]) == Decimal(50)
    assert Decimal(position["market_value"]) == Decimal(60)
    assert Decimal(position["unrealized_pl"]) == Decimal(10)
    assert Decimal(position["realized_pl"]) == Decimal(0)
    assert Decimal(position["dividends"]) == Decimal("0.25")
    assert Decimal(position["total_result"]) == Decimal("10.25")

    currency_summary = summary_response.json()["currencies"][0]
    assert currency_summary["currency"] == "USD"
    assert Decimal(currency_summary["total_result"]) == Decimal("10.25")
    assert currency_summary["priced_positions"] == 1
    assert currency_summary["total_open_positions"] == 1


def test_update_investment_price_rejects_invalid_or_unknown_ticker(
    clean_database: None,
) -> None:
    account_response = client.post(
        "/accounts",
        json={
            "name": "Investments",
            "bank_name": "Revolut",
            "currency": "EUR",
        },
    )
    account_id = account_response.json()["id"]

    invalid_price_response = client.put(
        f"/accounts/{account_id}/investment-prices/FAKE",
        json={"price": "0"},
    )
    unknown_ticker_response = client.put(
        f"/accounts/{account_id}/investment-prices/UNKNOWN",
        json={"price": "100"},
    )

    assert invalid_price_response.status_code == 422
    assert unknown_ticker_response.status_code == 404
    assert unknown_ticker_response.json() == {"detail": "Ticker not found"}


def test_refresh_investment_prices_updates_supported_positions(
    clean_database: None,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    account_response = client.post(
        "/accounts",
        json={
            "name": "Investments",
            "bank_name": "Revolut",
            "currency": "EUR",
        },
    )
    account_id = account_response.json()["id"]
    client.post(
        f"/accounts/{account_id}/imports/revolut-investments",
        files={
            "file": (
                "revolut-investments.csv",
                _create_fake_revolut_investment_csv(),
                "text/csv",
            )
        },
    )
    client.post(
        f"/accounts/{account_id}/imports/revolut-investments",
        files={
            "file": (
                "eur-investment.csv",
                (
                    b"Date,Ticker,Type,Quantity,Price per share,Total Amount,Currency,FX Rate\n"
                    b"2026-01-11T12:30:00Z,AIL,BUY - MARKET,1,EUR 172.46,EUR 172.46,EUR,1\n"
                ),
                "text/csv",
            )
        },
    )
    captured_call: tuple[dict[str, str], dict[str, str]] | None = None

    quoted_at = datetime.now(UTC).replace(microsecond=0) - timedelta(minutes=20)

    def fake_fetch(
        symbols: dict[str, str], currencies: dict[str, str]
    ) -> MarketPriceResult:
        nonlocal captured_call
        captured_call = (symbols, currencies)
        return MarketPriceResult(
            prices={
                "AIL": MarketQuote(Decimal("166.78"), quoted_at),
                "FAKE": MarketQuote(Decimal(120), quoted_at),
            },
            unavailable=[],
        )

    monkeypatch.setattr(investments_router, "fetch_yahoo_prices", fake_fetch)

    response = client.post(f"/accounts/{account_id}/investment-prices/refresh")

    assert response.status_code == 200
    assert response.json() == {
        "updated": ["AIL", "FAKE"],
        "unavailable": [],
        "manual_only": [],
    }
    assert captured_call == (
        {"AIL": "AI.PA", "FAKE": "FAKE"},
        {"AIL": "EUR", "FAKE": "USD"},
    )

    summary_response = client.get(
        "/investment-summary",
        params={"account_id": account_id},
    )
    current_prices = {
        position["ticker"]: Decimal(position["current_price"])
        for position in summary_response.json()["positions"]
    }
    for position in summary_response.json()["positions"]:
        assert position["price_source"] == "yahoo"
        assert datetime.fromisoformat(position["price_quoted_at"]) == quoted_at
        assert position["price_as_of"] == quoted_at.date().isoformat()
    assert current_prices == {
        "AIL": Decimal("166.78"),
        "FAKE": Decimal(120),
    }


def test_refresh_investment_prices_keeps_manual_price_on_provider_failure(
    clean_database: None,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    account_response = client.post(
        "/accounts",
        json={
            "name": "Investments",
            "bank_name": "Revolut",
            "currency": "EUR",
        },
    )
    account_id = account_response.json()["id"]
    client.post(
        f"/accounts/{account_id}/imports/revolut-investments",
        files={
            "file": (
                "revolut-investments.csv",
                _create_fake_revolut_investment_csv(),
                "text/csv",
            )
        },
    )

    assert (
        client.put(
            f"/accounts/{account_id}/investment-prices/FAKE", json={"price": "100"}
        ).status_code
        == 200
    )
    saved_summary = client.get(
        "/investment-summary", params={"account_id": account_id}
    ).json()

    def failed_fetch(*args: object) -> MarketPriceResult:
        raise MarketPriceError("Request failed")

    monkeypatch.setattr(investments_router, "fetch_yahoo_prices", failed_fetch)
    response = client.post(f"/accounts/{account_id}/investment-prices/refresh")

    assert response.status_code == 502
    assert response.json() == {"detail": "Market price provider unavailable"}
    assert (
        client.get("/investment-summary", params={"account_id": account_id}).json()
        == saved_summary
    )


def test_update_transaction_category_persists_category(
    clean_database: None,
) -> None:
    account_response = client.post(
        "/accounts",
        json={
            "name": "Main account",
            "bank_name": "Ibercaja",
            "currency": "EUR",
        },
    )
    assert account_response.status_code == 201
    account_id = account_response.json()["id"]

    with SessionFactory.begin() as session:
        transaction = Transaction(
            account_id=account_id,
            import_fingerprint="c" * 64,
            operation_date=date(2026, 8, 22),
            value_date=date(2026, 8, 22),
            amount=Decimal("-25.00"),
            balance_after=Decimal("1000.00"),
            bank_concept="CARD",
            description="FAKE RESTAURANT",
        )
        session.add(transaction)
        session.flush()
        transaction_id = transaction.id

    assert transaction_id is not None

    response = client.patch(
        f"/transactions/{transaction_id}/category",
        json={"category": "food"},
    )

    assert response.status_code == 200
    assert response.json()["category"] == "food"

    list_response = client.get("/transactions")

    assert list_response.status_code == 200
    assert list_response.json()[0]["category"] == "food"


def test_update_transaction_category_returns_404_when_transaction_does_not_exist(
    clean_database: None,
) -> None:
    response = client.patch(
        "/transactions/999999999/category",
        json={"category": "food"},
    )

    assert response.status_code == 404
    assert response.json() == {"detail": "Transaction not found"}


def test_update_transaction_category_rejects_empty_category(
    clean_database: None,
) -> None:
    response = client.patch(
        "/transactions/999999999/category",
        json={"category": ""},
    )

    assert response.status_code == 422


@pytest.mark.parametrize("saved_source", ["ibkr", "yahoo", "manual"])
def test_refresh_keeps_newer_saved_prices(
    ibkr_account: int,
    monkeypatch: pytest.MonkeyPatch,
    saved_source: str,
) -> None:
    response = client.post(
        f"/accounts/{ibkr_account}/imports/ibkr-investments",
        files={"file": ("report.xml", _create_fake_ibkr_xml("15"), "application/xml")},
    )
    assert response.status_code == 201
    quoted_at = datetime(2026, 10, 1, 12, tzinfo=UTC)

    def quote_fetch(*args: object) -> MarketPriceResult:
        return MarketPriceResult(
            prices={"FAKE": MarketQuote(Decimal(20), quoted_at)}, unavailable=[]
        )

    monkeypatch.setattr(investments_router, "fetch_yahoo_prices", quote_fetch)
    if saved_source in {"yahoo", "manual"}:
        refresh = client.post(f"/accounts/{ibkr_account}/investment-prices/refresh")
        assert refresh.status_code == 200
        assert refresh.json()["updated"] == ["FAKE"]
        position = client.get(
            "/investment-summary", params={"account_id": ibkr_account}
        ).json()["positions"][0]
        assert position["price_source"] == "yahoo"
        assert datetime.fromisoformat(position["price_quoted_at"]) == quoted_at
    if saved_source == "manual":
        manual = client.put(
            f"/accounts/{ibkr_account}/investment-prices/FAKE", json={"price": "25"}
        )
        assert manual.status_code == 200
        assert manual.json()["quoted_at"] is None

    saved_summary = client.get(
        "/investment-summary", params={"account_id": ibkr_account}
    ).json()
    quoted_at = datetime(2026, 9, 30, 12, tzinfo=UTC)
    refresh = client.post(f"/accounts/{ibkr_account}/investment-prices/refresh")
    assert refresh.status_code == 200
    assert refresh.json() == {"updated": [], "unavailable": ["FAKE"], "manual_only": []}
    assert (
        client.get("/investment-summary", params={"account_id": ibkr_account}).json()
        == saved_summary
    )


def test_newer_ibkr_report_clears_previous_yahoo_quote_timestamp(
    ibkr_account: int,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    client.post(
        f"/accounts/{ibkr_account}/imports/ibkr-investments",
        files={
            "file": (
                "report.xml",
                _create_fake_ibkr_xml("15", "20260929"),
                "application/xml",
            )
        },
    )
    monkeypatch.setattr(
        investments_router,
        "fetch_yahoo_prices",
        lambda *args: MarketPriceResult(
            prices={
                "FAKE": MarketQuote(Decimal(20), datetime(2026, 9, 30, 12, tzinfo=UTC))
            },
            unavailable=[],
        ),
    )
    refresh = client.post(f"/accounts/{ibkr_account}/investment-prices/refresh")
    assert refresh.json()["updated"] == ["FAKE"]
    response = client.post(
        f"/accounts/{ibkr_account}/imports/ibkr-investments",
        files={
            "file": (
                "report.xml",
                _create_fake_ibkr_xml("21", "20261001"),
                "application/xml",
            )
        },
    )
    assert response.status_code == 201
    assert response.json()["prices_updated"] == 1
    position = client.get(
        "/investment-summary", params={"account_id": ibkr_account}
    ).json()["positions"][0]
    assert position["price_source"] == "ibkr"
    assert position["price_quoted_at"] is None
    assert Decimal(position["current_price"]) == Decimal(21)


@pytest.fixture
def share_adjustment_account(clean_database: None) -> int:
    response = client.post(
        "/accounts",
        json={
            "name": "Test investments",
            "bank_name": "Test broker",
            "currency": "USD",
        },
    )
    assert response.status_code == 201
    account_id = response.json()["id"]
    imported = client.post(
        f"/accounts/{account_id}/imports/revolut-investments",
        files={
            "file": ("synthetic.csv", _create_fake_revolut_investment_csv(), "text/csv")
        },
    )
    assert imported.status_code == 201
    return account_id


def test_share_adjustment_is_recorded_once_and_survives_reimport(
    share_adjustment_account: int,
) -> None:
    account_id = share_adjustment_account
    response = client.put(
        f"/accounts/{account_id}/investment-share-adjustments/fake/2026-01-20",
        json={"multiplier": "1.1"},
    )
    assert response.status_code == 201
    assert response.json()["activity_type"] == "SHARE ADJUSTMENT"
    assert response.json()["quantity_multiplier"] == "1.10000000"
    assert response.json()["quantity"] is None
    assert Decimal(response.json()["total_amount"]) == 0
    repeated = client.put(
        f"/accounts/{account_id}/investment-share-adjustments/FAKE/2026-01-20",
        json={"multiplier": "1.10"},
    )
    assert repeated.status_code == 200
    assert repeated.json()["id"] == response.json()["id"]
    imported = client.post(
        f"/accounts/{account_id}/imports/revolut-investments",
        files={
            "file": ("synthetic.csv", _create_fake_revolut_investment_csv(), "text/csv")
        },
    )
    assert imported.json() == {"imported": 0}
    price = client.put(
        f"/accounts/{account_id}/investment-prices/FAKE",
        json={"price": "100"},
    )
    assert price.status_code == 200
    summary = client.get("/investment-summary", params={"account_id": account_id})
    assert summary.status_code == 200
    position = summary.json()["positions"][0]
    assert Decimal(position["quantity"]) == Decimal("0.55")
    assert Decimal(position["remaining_cost"]) == Decimal(50)
    assert Decimal(position["market_value"]) == Decimal(55)
    assert Decimal(position["unrealized_pl"]) == Decimal(5)
    activities = client.get(
        "/investment-activities", params={"account_id": account_id}
    ).json()
    assert len(activities) == 3
    purchase = next(
        activity
        for activity in activities
        if activity["activity_type"] == "BUY - MARKET"
    )
    assert Decimal(purchase["quantity"]) == Decimal("0.5")


def test_share_adjustment_rejects_conflicting_multiplier(
    share_adjustment_account: int,
) -> None:
    url = f"/accounts/{share_adjustment_account}/investment-share-adjustments/FAKE/2026-01-20"
    assert client.put(url, json={"multiplier": "1.1"}).status_code == 201
    response = client.put(url, json={"multiplier": "2"})
    assert response.status_code == 409
    assert response.json() == {
        "detail": "A different share adjustment is already recorded for this date"
    }
    activities = client.get(
        "/investment-activities", params={"account_id": share_adjustment_account}
    ).json()
    assert len(activities) == 3


@pytest.mark.parametrize(
    "multiplier", ["0", "-1", "1", "NaN", "Infinity", "0.000000001", "100000000"]
)
def test_share_adjustment_rejects_invalid_input(
    share_adjustment_account: int, multiplier: str
) -> None:
    response = client.put(
        f"/accounts/{share_adjustment_account}/investment-share-adjustments/FAKE/2026-01-20",
        json={"multiplier": multiplier},
    )
    assert response.status_code == 422
    activities = client.get(
        "/investment-activities", params={"account_id": share_adjustment_account}
    ).json()
    assert len(activities) == 2


@pytest.mark.parametrize(
    "effective_date",
    ["2026-01-01", (datetime.now(UTC).date() + timedelta(days=1)).isoformat()],
)
def test_share_adjustment_rejects_invalid_date(
    share_adjustment_account: int, effective_date: str
) -> None:
    response = client.put(
        f"/accounts/{share_adjustment_account}/investment-share-adjustments/FAKE/{effective_date}",
        json={"multiplier": "1.1"},
    )
    assert response.status_code == 422
    activities = client.get(
        "/investment-activities", params={"account_id": share_adjustment_account}
    ).json()
    assert len(activities) == 2


def test_share_adjustment_rejects_unknown_account_and_ticker(
    share_adjustment_account: int,
) -> None:
    for account_id, ticker in [(0, "FAKE"), (share_adjustment_account, "UNKNOWN")]:
        response = client.put(
            f"/accounts/{account_id}/investment-share-adjustments/{ticker}/2026-01-20",
            json={"multiplier": "1.1"},
        )
        assert response.status_code == 404


def test_share_adjustment_only_changes_selected_account(
    share_adjustment_account: int,
) -> None:
    other = client.post(
        "/accounts",
        json={
            "name": "Other investments",
            "bank_name": "Test broker",
            "currency": "USD",
        },
    ).json()["id"]
    imported = client.post(
        f"/accounts/{other}/imports/revolut-investments",
        files={
            "file": ("synthetic.csv", _create_fake_revolut_investment_csv(), "text/csv")
        },
    )
    assert imported.status_code == 201
    response = client.put(
        f"/accounts/{share_adjustment_account}/investment-share-adjustments/FAKE/2026-01-20",
        json={"multiplier": "1.1"},
    )
    assert response.status_code == 201
    summary = client.get("/investment-summary", params={"account_id": other}).json()
    assert Decimal(summary["positions"][0]["quantity"]) == Decimal("0.5")


@pytest.mark.parametrize(
    "overrides",
    [
        {"quantity_multiplier": None},
        {"quantity_multiplier": Decimal("NaN")},
        {"quantity": Decimal(1)},
        {"total_amount": Decimal(1)},
        {"activity_type": "BUY - MARKET"},
    ],
)
def test_database_rejects_malformed_share_adjustments(
    share_adjustment_account: int, overrides: dict
) -> None:
    values = {
        "account_id": share_adjustment_account,
        "import_fingerprint": "a" * 64,
        "occurred_at": datetime(2026, 1, 20, tzinfo=UTC),
        "ticker": "FAKE",
        "activity_type": "SHARE ADJUSTMENT",
        "quantity": None,
        "quantity_multiplier": Decimal("1.1"),
        "price_per_share": None,
        "total_amount": Decimal(0),
        "currency": "USD",
        "fx_rate": Decimal(1),
    }
    values.update(overrides)
    with pytest.raises(IntegrityError), SessionFactory.begin() as session:
        session.add(InvestmentActivity(**values))


def test_reverse_adjustment_rejects_history_with_too_many_shares_sold(
    share_adjustment_account: int,
) -> None:
    csv_contents = (
        b"Date,Ticker,Type,Quantity,Price per share,Total Amount,Currency,FX Rate\n"
        b"2026-01-25T12:00:00Z,FAKE,SELL - MARKET,0.4,USD 100,USD 40,USD,1\n"
    )
    imported = client.post(
        f"/accounts/{share_adjustment_account}/imports/revolut-investments",
        files={"file": ("synthetic.csv", csv_contents, "text/csv")},
    )
    assert imported.status_code == 201
    response = client.put(
        f"/accounts/{share_adjustment_account}/investment-share-adjustments/FAKE/2026-01-20",
        json={"multiplier": "0.5"},
    )
    assert response.status_code == 422
    summary = client.get(
        "/investment-summary", params={"account_id": share_adjustment_account}
    )
    assert summary.status_code == 200
    assert Decimal(summary.json()["positions"][0]["quantity"]) == Decimal("0.1")


def test_migration_cannot_discard_recorded_share_adjustments(
    share_adjustment_account: int,
) -> None:
    response = client.put(
        f"/accounts/{share_adjustment_account}/investment-share-adjustments/FAKE/2026-01-20",
        json={"multiplier": "1.1"},
    )
    assert response.status_code == 201
    with pytest.raises(RuntimeError, match="Cannot downgrade"):
        command.downgrade(Config("alembic.ini"), "a4c8e91d6b20")
    activities = client.get(
        "/investment-activities", params={"account_id": share_adjustment_account}
    ).json()
    assert len(activities) == 3
    assert (
        next(
            activity["quantity_multiplier"]
            for activity in activities
            if activity["activity_type"] == "SHARE ADJUSTMENT"
        )
        == "1.10000000"
    )


def test_delete_activity_keeps_price_until_last_ticker_activity(
    share_adjustment_account: int,
) -> None:
    account_id = share_adjustment_account
    assert (
        client.put(
            f"/accounts/{account_id}/investment-prices/FAKE", json={"price": "100"}
        ).status_code
        == 200
    )
    activities = client.get(
        "/investment-activities", params={"account_id": account_id}
    ).json()
    dividend = next(row for row in activities if row["activity_type"] == "DIVIDEND")
    purchase = next(row for row in activities if row["activity_type"] == "BUY - MARKET")
    deleted = client.delete(
        f"/accounts/{account_id}/investment-activities/{dividend['id']}"
    )
    assert deleted.status_code == 204 and deleted.content == b""
    position = client.get(
        "/investment-summary", params={"account_id": account_id}
    ).json()["positions"][0]
    assert Decimal(position["quantity"]) == Decimal("0.5")
    assert Decimal(position["dividends"]) == 0
    assert Decimal(position["current_price"]) == 100
    deleted = client.delete(
        f"/accounts/{account_id}/investment-activities/{purchase['id']}"
    )
    assert deleted.status_code == 204
    assert client.get(
        "/investment-summary", params={"account_id": account_id}
    ).json() == {"positions": [], "currencies": []}
    with SessionFactory() as session:
        assert (
            session.scalar(
                select(InvestmentPrice).where(InvestmentPrice.account_id == account_id)
            )
            is None
        )
        assert session.get(Account, account_id) is not None


@pytest.mark.parametrize("dependent_type", ["sale", "share adjustment"])
def test_delete_purchase_rejects_dependent_activity(
    share_adjustment_account: int, dependent_type: str
) -> None:
    account_id = share_adjustment_account
    activities = client.get(
        "/investment-activities", params={"account_id": account_id}
    ).json()
    purchase = next(row for row in activities if row["activity_type"] == "BUY - MARKET")
    if dependent_type == "sale":
        contents = b"Date,Ticker,Type,Quantity,Price per share,Total Amount,Currency,FX Rate\n2026-02-02T12:00:00Z,FAKE,SELL - MARKET,0.2,USD 100,USD 20,USD,1\n"
        result = client.post(
            f"/accounts/{account_id}/imports/revolut-investments",
            files={"file": ("synthetic.csv", contents, "text/csv")},
        )
        assert result.status_code == 201
        expected_quantity = Decimal("0.3")
    else:
        result = client.put(
            f"/accounts/{account_id}/investment-share-adjustments/FAKE/2026-01-20",
            json={"multiplier": "1.1"},
        )
        assert result.status_code == 201
        expected_quantity = Decimal("0.55")
    deleted = client.delete(
        f"/accounts/{account_id}/investment-activities/{purchase['id']}"
    )
    assert deleted.status_code == 409
    assert deleted.json() == {
        "detail": "Deletion would leave an invalid investment history"
    }
    assert (
        len(
            client.get(
                "/investment-activities", params={"account_id": account_id}
            ).json()
        )
        == 3
    )
    position = client.get(
        "/investment-summary", params={"account_id": account_id}
    ).json()["positions"][0]
    assert Decimal(position["quantity"]) == expected_quantity


def test_delete_activity_rejects_unknown_and_other_account_activity(
    share_adjustment_account: int,
) -> None:
    activities = client.get(
        "/investment-activities", params={"account_id": share_adjustment_account}
    ).json()
    other = client.post(
        "/accounts",
        json={"name": "Other account", "bank_name": "Test broker", "currency": "USD"},
    ).json()["id"]
    for account_id, activity_id in [
        (0, activities[0]["id"]),
        (share_adjustment_account, 0),
        (other, activities[0]["id"]),
    ]:
        response = client.delete(
            f"/accounts/{account_id}/investment-activities/{activity_id}"
        )
        assert response.status_code == 404
    assert (
        client.get(
            "/investment-activities", params={"account_id": share_adjustment_account}
        ).json()
        == activities
    )


def test_clear_investments_preserves_bank_data_and_other_accounts_and_allows_reimport(
    share_adjustment_account: int,
) -> None:
    account_id = share_adjustment_account
    other = client.post(
        "/accounts",
        json={"name": "Other account", "bank_name": "Test broker", "currency": "USD"},
    ).json()["id"]
    uploaded = client.post(
        f"/accounts/{other}/imports/revolut-investments",
        files={
            "file": ("synthetic.csv", _create_fake_revolut_investment_csv(), "text/csv")
        },
    )
    assert uploaded.status_code == 201
    assert (
        client.put(
            f"/accounts/{account_id}/investment-prices/FAKE", json={"price": "100"}
        ).status_code
        == 200
    )
    assert (
        client.put(
            f"/accounts/{other}/investment-prices/FAKE", json={"price": "200"}
        ).status_code
        == 200
    )
    assert (
        client.put(
            f"/accounts/{account_id}/investment-share-adjustments/FAKE/2026-01-20",
            json={"multiplier": "1.1"},
        ).status_code
        == 201
    )
    with SessionFactory.begin() as session:
        session.get(Account, account_id).current_balance = Decimal(123)
        session.add(
            Transaction(
                account_id=account_id,
                import_fingerprint="b" * 64,
                operation_date=date(2026, 1, 1),
                value_date=date(2026, 1, 1),
                amount=Decimal(123),
                balance_after=Decimal(123),
                bank_concept="TEST",
                description="SYNTHETIC BANK TRANSACTION",
            )
        )
    other_summary = client.get(
        "/investment-summary", params={"account_id": other}
    ).json()
    for _ in range(2):
        cleared = client.delete(f"/accounts/{account_id}/investments")
        assert cleared.status_code == 204 and cleared.content == b""
    assert (
        client.get("/investment-activities", params={"account_id": account_id}).json()
        == []
    )
    assert client.get(
        "/investment-summary", params={"account_id": account_id}
    ).json() == {"positions": [], "currencies": []}
    assert (
        client.get("/investment-summary", params={"account_id": other}).json()
        == other_summary
    )
    assert (
        len(client.get("/transactions", params={"account_id": account_id}).json()) == 1
    )
    with SessionFactory() as session:
        assert session.get(Account, account_id).current_balance == Decimal(123)
        assert (
            session.scalar(
                select(InvestmentPrice).where(InvestmentPrice.account_id == account_id)
            )
            is None
        )
    imported = client.post(
        f"/accounts/{account_id}/imports/revolut-investments",
        files={
            "file": ("synthetic.csv", _create_fake_revolut_investment_csv(), "text/csv")
        },
    )
    assert imported.json() == {"imported": 2}
    position = client.get(
        "/investment-summary", params={"account_id": account_id}
    ).json()["positions"][0]
    assert Decimal(position["quantity"]) == Decimal("0.5")


def test_clear_investments_rejects_unknown_account(
    share_adjustment_account: int,
) -> None:
    response = client.delete("/accounts/0/investments")
    assert response.status_code == 404
    assert (
        len(
            client.get(
                "/investment-activities",
                params={"account_id": share_adjustment_account},
            ).json()
        )
        == 2
    )
