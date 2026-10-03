from datetime import date
from decimal import Decimal
from pathlib import Path

import pytest
from openpyxl import Workbook

from transactions.ibercaja_importer import parse_ibercaja_xlsx


def test_parse_ibercaja_xlsx_returns_normalized_transaction(tmp_path: Path) -> None:
    workbook_path = tmp_path / "ibercaja.xlsx"
    workbook = Workbook()
    sheet = workbook.active

    sheet["A1"] = "FAKE ACCOUNT REPORT"
    sheet["A2"] = "FAKE GENERATION DATE"

    headers = [
        "Nº Orden",
        "Fecha Oper",
        "Fecha Valor",
        "Concepto",
        "Descripción",
        "Referencia",
        "Importe",
        "Saldo",
    ]
    for column, header in enumerate(headers, start=1):
        sheet.cell(row=6, column=column, value=header)

    fake_row = [
        1,
        "17-08-2026",
        "17-08-2026",
        "CARD",
        "FAKE SUPERMARKET",
        "FAKE123",
        -42.64,
        14897.84,
    ]
    for column, value in enumerate(fake_row, start=1):
        sheet.cell(row=7, column=column, value=value)

    workbook.save(workbook_path)
    workbook.close()

    transactions = parse_ibercaja_xlsx(workbook_path)

    assert transactions == [
        {
            "import_fingerprint": (
                "a80f7c370114ac1e7bba9879f362c2beaaf1765d4b2e8a5a44fb4d590f316c11"
            ),
            "operation_date": date(2026, 8, 17),
            "value_date": date(2026, 8, 17),
            "bank_concept": "CARD",
            "description": "FAKE SUPERMARKET",
            "amount": Decimal("-42.64"),
            "balance_after": Decimal("14897.84"),
        }
    ]


def test_parse_ibercaja_xlsx_rejects_missing_headers(tmp_path: Path) -> None:
    workbook_path = tmp_path / "not-ibercaja.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet["A1"] = "UNEXPECTED WORKBOOK"
    workbook.save(workbook_path)
    workbook.close()

    with pytest.raises(
        ValueError,
        match="Ibercaja transaction headers were not found",
    ):
        parse_ibercaja_xlsx(workbook_path)


def test_parse_ibercaja_xlsx_rejects_invalid_transaction_row(
    tmp_path: Path,
) -> None:
    workbook_path = tmp_path / "invalid-row.xlsx"
    workbook = Workbook()
    sheet = workbook.active

    headers = [
        "Nº Orden",
        "Fecha Oper",
        "Fecha Valor",
        "Concepto",
        "Descripción",
        "Referencia",
        "Importe",
        "Saldo",
    ]
    for column, header in enumerate(headers, start=1):
        sheet.cell(row=6, column=column, value=header)

    invalid_row = [
        1,
        None,
        "17-08-2026",
        "CARD",
        "FAKE SUPERMARKET",
        "FAKE123",
        -42.64,
        14897.84,
    ]
    for column, value in enumerate(invalid_row, start=1):
        sheet.cell(row=7, column=column, value=value)

    workbook.save(workbook_path)
    workbook.close()

    with pytest.raises(
        ValueError,
        match="Invalid Ibercaja transaction row 7",
    ):
        parse_ibercaja_xlsx(workbook_path)


def test_parse_ibercaja_xlsx_rejects_invalid_amount(tmp_path: Path) -> None:
    workbook_path = tmp_path / "invalid-amount.xlsx"
    workbook = Workbook()
    sheet = workbook.active

    headers = [
        "Nº Orden",
        "Fecha Oper",
        "Fecha Valor",
        "Concepto",
        "Descripción",
        "Referencia",
        "Importe",
        "Saldo",
    ]
    for column, header in enumerate(headers, start=1):
        sheet.cell(row=6, column=column, value=header)

    invalid_row = [
        1,
        "17-08-2026",
        "17-08-2026",
        "CARD",
        "FAKE SUPERMARKET",
        "FAKE123",
        "NOT A NUMBER",
        14897.84,
    ]
    for column, value in enumerate(invalid_row, start=1):
        sheet.cell(row=7, column=column, value=value)

    workbook.save(workbook_path)
    workbook.close()

    with pytest.raises(
        ValueError,
        match="Invalid Ibercaja transaction row 7",
    ):
        parse_ibercaja_xlsx(workbook_path)


def test_parse_ibercaja_xlsx_rejects_missing_description(tmp_path: Path) -> None:
    workbook_path = tmp_path / "missing-description.xlsx"
    workbook = Workbook()
    sheet = workbook.active

    headers = [
        "Nº Orden",
        "Fecha Oper",
        "Fecha Valor",
        "Concepto",
        "Descripción",
        "Referencia",
        "Importe",
        "Saldo",
    ]
    for column, header in enumerate(headers, start=1):
        sheet.cell(row=6, column=column, value=header)

    invalid_row = [
        1,
        "17-08-2026",
        "17-08-2026",
        "CARD",
        None,
        "FAKE123",
        -42.64,
        14897.84,
    ]
    for column, value in enumerate(invalid_row, start=1):
        sheet.cell(row=7, column=column, value=value)

    workbook.save(workbook_path)
    workbook.close()

    with pytest.raises(
        ValueError,
        match="Invalid Ibercaja transaction row 7",
    ):
        parse_ibercaja_xlsx(workbook_path)
